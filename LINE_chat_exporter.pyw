import csv
import os
import re
import subprocess
import sys
import threading
from collections import Counter
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill


def app_base_dir() -> Path:
    """Directory for CSV/XLSX next to the .exe when frozen, else this script's folder."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


INDEX_PATTERN = "LINE_聊天室索引*.csv"
MESSAGES_FILE = "LINE_全部聊天.csv"
INDEX_REQUIRED = {"聊天室名稱", "聊天室ID", "類型", "訊息數"}
MESSAGE_REQUIRED = {
    "聊天室名稱",
    "聊天室ID",
    "時間（台北）",
    "方向",
    "發言者",
    "發言者ID",
    "訊息類型",
    "內容",
    "訊息ID",
}
INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
MESSAGE_HEADERS = [
    "時間（台北）",
    "方向",
    "發言者",
    "發言者ID",
    "訊息類型",
    "內容",
    "訊息ID",
]


def find_input_files(base_dir):
    candidates = sorted(base_dir.glob(INDEX_PATTERN))
    if not candidates:
        candidates = sorted(base_dir.parent.glob(f"**/{INDEX_PATTERN}"))
    if not candidates:
        return None, None
    index_path = candidates[0]
    messages_path = index_path.parent / MESSAGES_FILE
    return index_path, messages_path if messages_path.exists() else None


def read_csv_headers(path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        return set(next(reader, []))


def read_index(path):
    if not INDEX_REQUIRED.issubset(read_csv_headers(path)):
        raise ValueError("聊天室索引欄位不完整。")
    records = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            try:
                count = int(row.get("訊息數", 0) or 0)
            except ValueError:
                count = 0
            records.append(
                {
                    "name": row.get("聊天室名稱", "") or row.get("聊天室ID", ""),
                    "id": row.get("聊天室ID", ""),
                    "category": row.get("類別", "") or fallback_category(row.get("聊天室ID", "")),
                    "type": row.get("類型", ""),
                    "count": count,
                    "oldest": row.get("最早訊息（台北）", ""),
                    "newest": row.get("最新訊息（台北）", ""),
                }
            )
    return records


def fallback_category(chat_id):
    if chat_id.startswith("u"):
        return "好友"
    if chat_id.startswith("c"):
        return "群組"
    if chat_id.startswith(("m", "t")):
        return "社群"
    return "其他"


def unique_sheet_name(label, used_names):
    base = INVALID_SHEET_CHARS.sub("_", (label or "聊天室")).strip(" '")
    base = base or "聊天室"
    base = base[:31]
    candidate = base
    suffix = 2
    folded = {name.casefold() for name in used_names}
    while candidate.casefold() in folded:
        tail = f"_{suffix}"
        candidate = f"{base[: 31 - len(tail)]}{tail}"
        suffix += 1
    used_names.add(candidate)
    return candidate


def parse_time(value):
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d %H:%M:%S")
    except ValueError:
        return value


def styled_header_cell(sheet, value):
    cell = WriteOnlyCell(sheet, value=value)
    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    return cell


def data_cell(sheet, value, *, wrap=False, date=False):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        value = "'" + value
    cell = WriteOnlyCell(sheet, value=value)
    cell.font = Font(name="Arial", size=10, color="222222")
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    if date and isinstance(value, datetime):
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    return cell


def export_workbook(index_records, messages_path, selected_ids, output_path, progress=None):
    if not MESSAGE_REQUIRED.issubset(read_csv_headers(messages_path)):
        raise ValueError("全部聊天檔案欄位不完整。")

    selected = {record["id"]: record for record in index_records if record["id"] in selected_ids}
    if not selected:
        raise ValueError("尚未選取聊天室。")

    workbook = Workbook(write_only=True)
    used_names = {"聊天室索引"}
    index_sheet = workbook.create_sheet("聊天室索引")
    index_sheet.freeze_panes = "A2"
    index_sheet.column_dimensions["A"].width = 34
    index_sheet.column_dimensions["B"].width = 38
    index_sheet.column_dimensions["C"].width = 12
    index_sheet.column_dimensions["D"].width = 14
    index_sheet.column_dimensions["E"].width = 14
    index_sheet.column_dimensions["F"].width = 21
    index_sheet.column_dimensions["G"].width = 21
    index_sheet.append(
        [styled_header_cell(index_sheet, value) for value in [
            "聊天室名稱", "聊天室ID", "類別", "類型", "訊息數", "最早訊息", "最新訊息"
        ]]
    )

    sheets = {}
    counts = Counter()
    for record in selected.values():
        sheet_name = unique_sheet_name(record["name"], used_names)
        sheet = workbook.create_sheet(sheet_name)
        sheet.freeze_panes = "A2"
        widths = [21, 10, 24, 38, 17, 80, 38]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width
        sheet.append([styled_header_cell(sheet, value) for value in MESSAGE_HEADERS])
        sheets[record["id"]] = (sheet, sheet_name)
        index_sheet.append(
            [
                data_cell(index_sheet, record["name"]),
                data_cell(index_sheet, record["id"]),
                data_cell(index_sheet, record["category"]),
                data_cell(index_sheet, record["type"]),
                data_cell(index_sheet, record["count"]),
                data_cell(index_sheet, parse_time(record["oldest"]), date=True),
                data_cell(index_sheet, parse_time(record["newest"]), date=True),
            ]
        )

    if progress:
        progress("正在讀取全部聊天資料…")

    with messages_path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            chat_id = row.get("聊天室ID", "")
            target = sheets.get(chat_id)
            if target is None:
                continue
            sheet, _ = target
            created = parse_time(row.get("時間（台北）", ""))
            sheet.append(
                [
                    data_cell(sheet, created, date=True),
                    data_cell(sheet, row.get("方向", "")),
                    data_cell(sheet, row.get("發言者", "")),
                    data_cell(sheet, row.get("發言者ID", "")),
                    data_cell(sheet, row.get("訊息類型", "")),
                    data_cell(sheet, row.get("內容", ""), wrap=True),
                    data_cell(sheet, row.get("訊息ID", "")),
                ]
            )
            counts[chat_id] += 1

    for chat_id, (sheet, _) in sheets.items():
        sheet.auto_filter.ref = f"A1:G{counts[chat_id] + 1}"
    index_sheet.auto_filter.ref = f"A1:G{len(selected) + 1}"

    if progress:
        progress("正在建立 Excel 活頁簿…")
    workbook.save(output_path)
    return counts


class LineChatExporter(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("LINE 聊天匯出工具")
        self.geometry("1000x680")
        self.minsize(820, 520)
        self.records = []
        self.selected_ids = set()
        self.index_path = None
        self.messages_path = None
        self.search_var = tk.StringVar()
        self.category_vars = {
            name: tk.BooleanVar(value=True)
            for name in ("官方帳號", "社群", "群組", "好友")
        }
        self.status_var = tk.StringVar(value="正在搜尋 LINE_聊天室索引…")
        self._build_ui()
        self.after(100, self.auto_load)

    def _build_ui(self):
        container = ttk.Frame(self, padding=14)
        container.pack(fill="both", expand=True)

        title = ttk.Label(container, text="LINE 聊天匯出工具", font=("Microsoft JhengHei UI", 16, "bold"))
        title.pack(anchor="w")

        file_row = ttk.Frame(container)
        file_row.pack(fill="x", pady=(12, 8))
        self.file_label = ttk.Label(file_row, text="尚未載入聊天室索引")
        self.file_label.pack(side="left", fill="x", expand=True)
        self.update_button = ttk.Button(
            file_row,
            text="更新 LINE 資料與聯絡人名稱",
            command=self.start_update,
        )
        self.update_button.pack(side="right", padx=(8, 0))
        ttk.Button(file_row, text="重新選擇索引", command=self.choose_index).pack(side="right")

        search_row = ttk.Frame(container)
        search_row.pack(fill="x", pady=(0, 8))
        ttk.Label(search_row, text="搜尋：").pack(side="left")
        search_entry = ttk.Entry(search_row, textvariable=self.search_var)
        search_entry.pack(side="left", fill="x", expand=True)
        self.search_var.trace_add("write", lambda *_: self.refresh_tree())
        ttk.Button(search_row, text="清除搜尋", command=lambda: self.search_var.set("")).pack(side="left", padx=(8, 0))

        category_row = ttk.Frame(container)
        category_row.pack(fill="x", pady=(0, 8))
        ttk.Label(category_row, text="類別：").pack(side="left")
        for name, variable in self.category_vars.items():
            ttk.Checkbutton(
                category_row,
                text=name,
                variable=variable,
                command=self.refresh_tree,
            ).pack(side="left")
        ttk.Button(category_row, text="全選", width=5, command=lambda: self.set_categories(True)).pack(side="left", padx=(8, 0))
        ttk.Button(category_row, text="清除", width=5, command=lambda: self.set_categories(False)).pack(side="left", padx=(4, 0))

        columns = ("checked", "name", "category", "count", "oldest", "newest", "id")
        self.tree = ttk.Treeview(container, columns=columns, show="headings", selectmode="none")
        headings = {
            "checked": "勾選",
            "name": "聊天室名稱",
            "category": "類別",
            "count": "訊息數",
            "oldest": "最早訊息",
            "newest": "最新訊息",
            "id": "聊天室ID",
        }
        widths = {"checked": 60, "name": 240, "category": 90, "count": 90, "oldest": 150, "newest": 150, "id": 240}
        for column in columns:
            self.tree.heading(column, text=headings[column])
            self.tree.column(column, width=widths[column], anchor="center" if column in ("checked", "category", "count") else "w")
        self.tree.bind("<Button-1>", self.toggle_check)

        scrollbar = ttk.Scrollbar(container, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="left", fill="y")

        bottom = ttk.Frame(self, padding=(14, 0, 14, 14))
        bottom.pack(fill="x")
        ttk.Button(bottom, text="勾選目前搜尋結果", command=self.select_visible).pack(side="left")
        ttk.Button(bottom, text="取消全部", command=self.clear_selection).pack(side="left", padx=(8, 0))
        self.export_button = ttk.Button(bottom, text="匯出已勾選聊天室", command=self.start_export)
        self.export_button.pack(side="right")
        ttk.Label(bottom, textvariable=self.status_var).pack(side="right", padx=(0, 14))

    def auto_load(self):
        index_path, messages_path = find_input_files(app_base_dir())
        if index_path and messages_path:
            self.load_files(index_path, messages_path)
        else:
            self.status_var.set("找不到完整資料，請選擇 LINE_聊天室索引.csv")
            self.choose_index()

    def choose_index(self):
        path = filedialog.askopenfilename(
            title="選擇 LINE_聊天室索引.csv",
            filetypes=[("CSV 檔案", "*.csv")],
        )
        if not path:
            return
        index_path = Path(path)
        messages_path = index_path.parent / MESSAGES_FILE
        if not messages_path.exists():
            selected = filedialog.askopenfilename(
                title="選擇 LINE_全部聊天.csv",
                initialdir=str(index_path.parent),
                filetypes=[("CSV 檔案", "*.csv")],
            )
            if not selected:
                return
            messages_path = Path(selected)
        self.load_files(index_path, messages_path)

    def load_files(self, index_path, messages_path):
        try:
            records = read_index(index_path)
            if not MESSAGE_REQUIRED.issubset(read_csv_headers(messages_path)):
                raise ValueError("LINE_全部聊天.csv 欄位不完整。")
        except Exception as exc:
            messagebox.showerror("無法載入", str(exc))
            return
        self.index_path = index_path
        self.messages_path = messages_path
        self.records = records
        self.selected_ids.clear()
        self.file_label.config(text=f"已載入：{index_path.name}（{len(records)} 個聊天室）")
        self.status_var.set(f"已載入 {len(records)} 個聊天室，已勾選 0 個")
        self.refresh_tree()

    def filtered_records(self):
        keyword = self.search_var.get().strip().casefold()
        categories = {name for name, variable in self.category_vars.items() if variable.get()}
        return [
            record for record in self.records
            if record["category"] in categories
            and (not keyword or keyword in record["name"].casefold() or keyword in record["id"].casefold())
        ]

    def set_categories(self, selected):
        for variable in self.category_vars.values():
            variable.set(selected)
        self.refresh_tree()

    def refresh_tree(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for record in self.filtered_records():
            checked = "☑" if record["id"] in self.selected_ids else "☐"
            self.tree.insert(
                "",
                "end",
                iid=record["id"],
                values=(checked, record["name"], record["category"], f'{record["count"]:,}', record["oldest"], record["newest"], record["id"]),
            )

    def toggle_check(self, event):
        if self.tree.identify_region(event.x, event.y) != "cell":
            return
        item = self.tree.identify_row(event.y)
        column = self.tree.identify_column(event.x)
        if not item or column != "#1":
            return
        if item in self.selected_ids:
            self.selected_ids.remove(item)
        else:
            self.selected_ids.add(item)
        self.update_selection_status()
        self.refresh_tree()

    def select_visible(self):
        self.selected_ids.update(record["id"] for record in self.filtered_records())
        self.update_selection_status()
        self.refresh_tree()

    def clear_selection(self):
        self.selected_ids.clear()
        self.update_selection_status()
        self.refresh_tree()

    def update_selection_status(self):
        self.status_var.set(f"已載入 {len(self.records)} 個聊天室，已勾選 {len(self.selected_ids)} 個")

    def start_update(self):
        helper = app_base_dir() / "line_read_helper.py"
        if not getattr(sys, "frozen", False) and not helper.exists():
            messagebox.showerror("無法更新", f"找不到更新程式：\n{helper}")
            return
        self.update_button.config(state="disabled")
        self.export_button.config(state="disabled")
        self.status_var.set("正在讀取 LINE 並更新聯絡人名稱…")
        thread = threading.Thread(target=self._update_worker, args=(helper,), daemon=True)
        thread.start()

    def _update_worker(self, helper):
        try:
            base = app_base_dir()
            if getattr(sys, "frozen", False):
                import line_read_helper

                old_argv = sys.argv
                try:
                    sys.argv = [sys.argv[0], "--export-all"]
                    line_read_helper.main()
                finally:
                    sys.argv = old_argv
            else:
                local_vendor = base / "vendor"
                vendor = local_vendor if local_vendor.exists() else base.parent.parent / "work" / "vendor"
                environment = os.environ.copy()
                if vendor.exists():
                    existing = environment.get("PYTHONPATH", "")
                    environment["PYTHONPATH"] = str(vendor) + (os.pathsep + existing if existing else "")
                result = subprocess.run(
                    [sys.executable, str(helper), "--export-all"],
                    cwd=str(base),
                    env=environment,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
                )
                if result.returncode != 0:
                    details = (result.stdout + "\n" + result.stderr).strip()
                    raise RuntimeError(details or "LINE 資料更新失敗。")
            index_path = base / "LINE_聊天室索引.csv"
            messages_path = base / MESSAGES_FILE
            self.after(0, self._update_success, index_path, messages_path)
        except Exception as exc:
            self.after(0, self._update_failure, str(exc))

    def _update_success(self, index_path, messages_path):
        self.update_button.config(state="normal")
        self.export_button.config(state="normal")
        self.load_files(index_path, messages_path)
        messagebox.showinfo("更新完成", "聊天室索引與聯絡人名稱已更新。")

    def _update_failure(self, error):
        self.update_button.config(state="normal")
        self.export_button.config(state="normal")
        self.update_selection_status()
        messagebox.showerror(
            "更新失敗",
            f"請確認 LINE 電腦版已登入，並開啟任一聊天室後再試一次。\n\n{error}",
        )

    def start_export(self):
        if not self.selected_ids:
            messagebox.showinfo("尚未勾選", "請先勾選至少一個聊天室。")
            return
        default_name = f"LINE_聊天匯出_{datetime.now():%Y%m%d_%H%M}.xlsx"
        output = filedialog.asksaveasfilename(
            title="儲存 Excel 活頁簿",
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[("Excel 活頁簿", "*.xlsx")],
        )
        if not output:
            return
        self.export_button.config(state="disabled")
        selected = set(self.selected_ids)
        thread = threading.Thread(target=self._export_worker, args=(selected, Path(output)), daemon=True)
        thread.start()

    def _export_worker(self, selected, output_path):
        try:
            counts = export_workbook(
                self.records,
                self.messages_path,
                selected,
                output_path,
                progress=lambda text: self.after(0, self.status_var.set, text),
            )
            total = sum(counts.values())
            self.after(0, self._export_success, output_path, len(selected), total)
        except Exception as exc:
            self.after(0, self._export_failure, str(exc))

    def _export_success(self, output_path, chat_count, message_count):
        self.export_button.config(state="normal")
        self.update_selection_status()
        messagebox.showinfo(
            "匯出完成",
            f"已匯出 {chat_count} 個聊天室、{message_count:,} 則訊息。\n\n{output_path}",
        )

    def _export_failure(self, error):
        self.export_button.config(state="normal")
        self.update_selection_status()
        messagebox.showerror("匯出失敗", error)


def self_test():
    base = app_base_dir()
    index_path, messages_path = find_input_files(base)
    if not index_path or not messages_path:
        raise RuntimeError("找不到測試所需的索引或聊天 CSV。")
    records = read_index(index_path)
    selected = {record["id"] for record in records if record["count"] > 0}
    selected = set(list(selected)[:2])
    output = base.parent / "work" / "line_exporter_self_test.xlsx"
    counts = export_workbook(records, messages_path, selected, output)
    workbook = load_workbook(output, read_only=True)
    try:
        assert workbook.sheetnames[0] == "聊天室索引"
        assert len(workbook.sheetnames) == len(selected) + 1
        assert sum(counts.values()) > 0
        assert sum(1 for _ in workbook["聊天室索引"].rows) == len(selected) + 1
        assert sum(sum(1 for _ in sheet.rows) - 1 for sheet in workbook.worksheets[1:]) == sum(counts.values())
        for name in workbook.sheetnames:
            assert len(name) <= 31
            assert not INVALID_SHEET_CHARS.search(name)
    finally:
        workbook.close()
    print(f"SELF_TEST_OK sheets={len(selected)} messages={sum(counts.values())} output={output}")


if __name__ == "__main__":
    if "--self-test" in sys.argv:
        self_test()
    else:
        LineChatExporter().mainloop()
